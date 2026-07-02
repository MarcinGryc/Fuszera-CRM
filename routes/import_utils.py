from openpyxl import load_workbook


def read_xlsx_rows(file):
    workbook = load_workbook(file, data_only=True)
    sheet = workbook.active

    rows = list(sheet.iter_rows(values_only=True))

    if not rows:
        return []

    headers = [str(header).strip() if header is not None else "" for header in rows[0]]

    data = []

    for row in rows[1:]:
        item = {}

        for index, header in enumerate(headers):
            if not header:
                continue

            item[header] = row[index] if index < len(row) else None

        if any(value not in (None, "") for value in item.values()):
            data.append(item)

    return data