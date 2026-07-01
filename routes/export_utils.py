from io import BytesIO
from flask import send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def export_xlsx(filename, sheets):
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    header_fill = PatternFill("solid", fgColor="FFB6D9")

    for sheet_name, headers, rows in sheets:
        ws = wb.create_sheet(sheet_name[:31])
        ws.append(headers)

        for cell in ws[1]:
            cell.font = Font(bold=True, color="1F1F24")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for row in rows:
            ws.append(row)

        for column_cells in ws.columns:
            max_length = 0
            column_letter = column_cells[0].column_letter

            for cell in column_cells:
                value = str(cell.value) if cell.value is not None else ""
                max_length = max(max_length, len(value))

            ws.column_dimensions[column_letter].width = min(max_length + 3, 45)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )